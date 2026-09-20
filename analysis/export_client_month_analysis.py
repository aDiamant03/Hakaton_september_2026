from __future__ import annotations

import csv
import json
import shutil
from datetime import datetime
from pathlib import Path

import pandas as pd
import xlsxwriter
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / "working_data" / "client_month_analysis"
OUT = ROOT / "outputs" / "client_month_analysis"
FIG = ROOT / "figures" / "client_month_analysis"
TMP = ROOT / "tmp" / "client_month_analysis"
for directory in (OUT, TMP):
    directory.mkdir(parents=True, exist_ok=True)

OUTPUT = OUT / "DANO_клиент_месяц_анализ.xlsx"
OUTPUT_CSV = OUT / "DANO_клиент_месяц_общая.csv"
PREVIEW = TMP / "preview.xlsx"


def add_formats(workbook: xlsxwriter.Workbook) -> dict[str, xlsxwriter.format.Format]:
    return {
        "title": workbook.add_format({"font_name": "Arial", "font_size": 24, "bold": True, "font_color": "#111827"}),
        "subtitle": workbook.add_format({"font_name": "Arial", "font_size": 11, "font_color": "#475569", "text_wrap": True, "valign": "top"}),
        "section": workbook.add_format({"font_name": "Arial", "font_size": 14, "bold": True, "font_color": "#111827", "bg_color": "#FEF3C7", "bottom": 2, "bottom_color": "#F59E0B"}),
        "kpi_label": workbook.add_format({"font_name": "Arial", "font_size": 10, "font_color": "#475569", "bg_color": "#FFFFFF", "align": "center", "valign": "vcenter", "border": 1, "border_color": "#E2E8F0"}),
        "kpi_value": workbook.add_format({"font_name": "Arial", "font_size": 18, "bold": True, "font_color": "#0F766E", "bg_color": "#FFFFFF", "align": "center", "valign": "vcenter", "border": 1, "border_color": "#E2E8F0", "num_format": "#,##0"}),
        "kpi_pct": workbook.add_format({"font_name": "Arial", "font_size": 18, "bold": True, "font_color": "#0F766E", "bg_color": "#FFFFFF", "align": "center", "valign": "vcenter", "border": 1, "border_color": "#E2E8F0", "num_format": "0.0%"}),
        "header": workbook.add_format({"font_name": "Arial", "font_size": 9, "bold": True, "font_color": "#FFFFFF", "bg_color": "#334155", "align": "center", "valign": "vcenter", "text_wrap": True, "border": 1, "border_color": "#FFFFFF"}),
        "header_demo": workbook.add_format({"font_name": "Arial", "font_size": 9, "bold": True, "font_color": "#FFFFFF", "bg_color": "#374151", "align": "center", "valign": "vcenter", "text_wrap": True, "border": 1, "border_color": "#FFFFFF"}),
        "header_fines": workbook.add_format({"font_name": "Arial", "font_size": 9, "bold": True, "font_color": "#FFFFFF", "bg_color": "#B45309", "align": "center", "valign": "vcenter", "text_wrap": True, "border": 1, "border_color": "#FFFFFF"}),
        "header_fuel": workbook.add_format({"font_name": "Arial", "font_size": 9, "bold": True, "font_color": "#FFFFFF", "bg_color": "#0F766E", "align": "center", "valign": "vcenter", "text_wrap": True, "border": 1, "border_color": "#FFFFFF"}),
        "text": workbook.add_format({"font_name": "Arial", "font_size": 9, "font_color": "#1F2937", "valign": "vcenter"}),
        "text_wrap": workbook.add_format({"font_name": "Arial", "font_size": 9, "font_color": "#1F2937", "valign": "top", "text_wrap": True}),
        "integer": workbook.add_format({"font_name": "Arial", "font_size": 9, "num_format": "#,##0", "valign": "vcenter"}),
        "decimal": workbook.add_format({"font_name": "Arial", "font_size": 9, "num_format": "#,##0.00", "valign": "vcenter"}),
        "money": workbook.add_format({"font_name": "Arial", "font_size": 9, "num_format": "#,##0.00 [$₽-ru-RU]", "valign": "vcenter"}),
        "percent": workbook.add_format({"font_name": "Arial", "font_size": 9, "num_format": "0.0%", "valign": "vcenter"}),
        "date": workbook.add_format({"font_name": "Arial", "font_size": 9, "num_format": "yyyy-mm-dd", "valign": "vcenter"}),
        "note": workbook.add_format({"font_name": "Arial", "font_size": 10, "font_color": "#475569", "bg_color": "#F8FAFC", "text_wrap": True, "valign": "top", "border": 1, "border_color": "#E2E8F0"}),
        "method_head": workbook.add_format({"font_name": "Arial", "font_size": 11, "bold": True, "font_color": "#FFFFFF", "bg_color": "#0F766E", "text_wrap": True, "valign": "vcenter", "border": 1, "border_color": "#FFFFFF"}),
        "matrix_header": workbook.add_format({"font_name": "Arial", "font_size": 8, "bold": True, "font_color": "#FFFFFF", "bg_color": "#334155", "align": "center", "valign": "vcenter", "text_wrap": True, "border": 1, "border_color": "#FFFFFF"}),
        "matrix": workbook.add_format({"font_name": "Arial", "font_size": 8, "num_format": "0.00", "align": "center", "valign": "vcenter"}),
    }


def setup_sheet(sheet, zoom=90):
    sheet.hide_gridlines(2)
    sheet.set_zoom(zoom)
    sheet.set_default_row(17)


def write_frame(sheet, frame: pd.DataFrame, formats: dict, start_row=0, start_col=0, autofilter=True):
    for j, header in enumerate(frame.columns):
        sheet.write(start_row, start_col + j, header, formats["header"])
    sheet.set_row(start_row, 38)
    for i, row in enumerate(frame.itertuples(index=False, name=None), start=start_row + 1):
        for j, value in enumerate(row):
            col = frame.columns[j]
            if pd.isna(value):
                sheet.write_blank(i, start_col + j, None, formats["text"])
            elif isinstance(value, pd.Timestamp):
                sheet.write_datetime(i, start_col + j, value.to_pydatetime(), formats["date"])
            elif isinstance(value, (int, float)) and not isinstance(value, bool):
                if "pct" in col or "share" in col:
                    sheet.write_number(i, start_col + j, float(value), formats["percent"])
                elif any(token in col for token in ("price", "amount", "spend")):
                    sheet.write_number(i, start_col + j, float(value), formats["money"])
                elif any(token in col for token in ("volume", "per_client", "median", "mean", "weighted")):
                    sheet.write_number(i, start_col + j, float(value), formats["decimal"])
                else:
                    sheet.write_number(i, start_col + j, float(value), formats["integer"])
            else:
                sheet.write(i, start_col + j, str(value), formats["text"])
    if autofilter and len(frame):
        sheet.autofilter(start_row, start_col, start_row + len(frame), start_col + len(frame.columns) - 1)
    sheet.freeze_panes(start_row + 1, start_col)


def add_chart(workbook, sheet, chart_type, title, categories, series, position, y_name=None):
    chart = workbook.add_chart({"type": chart_type})
    for item in series:
        chart.add_series(item)
    chart.set_title({"name": title, "name_font": {"name": "Arial", "size": 12, "bold": True}})
    chart.set_x_axis({"name": "Месяц", "label_position": "low", "name_font": {"name": "Arial", "size": 9}, "num_font": {"name": "Arial", "size": 8}})
    chart.set_y_axis({"name": y_name or "", "major_gridlines": {"visible": True, "line": {"color": "#E2E8F0"}}, "name_font": {"name": "Arial", "size": 9}, "num_font": {"name": "Arial", "size": 8}})
    chart.set_legend({"position": "bottom", "font": {"name": "Arial", "size": 8}})
    chart.set_plotarea({"fill": {"color": "#FFFFFF"}, "border": {"none": True}})
    chart.set_chartarea({"fill": {"color": "#FFFFFF"}, "border": {"color": "#E2E8F0"}})
    chart.set_style(10)
    sheet.insert_chart(position, chart, {"x_scale": 1.18, "y_scale": 1.18})


def build_workbook(path: Path, include_panel: bool = True):
    metadata = json.loads((WORK / "metadata.json").read_text(encoding="utf-8"))
    monthly = pd.read_csv(WORK / "monthly_summary_quality.csv", parse_dates=["month"])
    regions = pd.read_csv(WORK / "region_summary_quality.csv")
    quality = pd.read_csv(WORK / "fuel_quality_summary.csv")
    histograms = pd.read_csv(WORK / "fuel_histograms.csv")
    corr_level = pd.read_csv(WORK / "correlation_level_spearman.csv", index_col=0)
    corr_change = pd.read_csv(WORK / "correlation_change_spearman.csv", index_col=0)

    workbook = xlsxwriter.Workbook(path, {
        "constant_memory": include_panel,
        "strings_to_numbers": False,
        "strings_to_formulas": False,
        "strings_to_urls": False,
        "nan_inf_to_errors": True,
    })
    workbook.set_properties({
        "title": "DANO — клиент-месяц, выбросы, графики и карты",
        "subject": "Агрегация по client_id × месяц, апрель–август 2026",
        "company": "DATA SWAGA team",
    })
    fmt = add_formats(workbook)
    overview = workbook.add_worksheet("Обзор")
    monthly_sheet = workbook.add_worksheet("Месяцы")
    region_sheet = workbook.add_worksheet("Регионы")
    quality_sheet = workbook.add_worksheet("Выбросы")
    corr_sheet = workbook.add_worksheet("Корреляции")
    change_corr_sheet = workbook.add_worksheet("Корр_изменений")
    maps_sheet = workbook.add_worksheet("Карты")
    methods_sheet = workbook.add_worksheet("Методика")
    panel_sheet = workbook.add_worksheet("Клиент-месяц") if include_panel else None

    for sheet in [overview, monthly_sheet, region_sheet, quality_sheet, corr_sheet, change_corr_sheet, maps_sheet, methods_sheet] + ([panel_sheet] if panel_sheet else []):
        setup_sheet(sheet, 85 if sheet is panel_sheet else 90)

    # Overview/dashboard.
    overview.set_tab_color("#FACC15")
    overview.set_column("A:A", 3)
    overview.set_column("B:Q", 12)
    overview.set_row(1, 36)
    overview.merge_range("B2:Q2", "DANO — клиент-месяц: штрафы, топливо, выбросы и регионы", fmt["title"])
    overview.merge_range(
        "B3:Q4",
        "Панель построена на едином ключе client_id × месяц. Raw-показатели сохранены; для основного анализа топлива используется отдельный робастный сценарий: положительный объём и цена 50–130 ₽/л.",
        fmt["subtitle"],
    )
    kpis = [
        ("B6:D6", "B7:D8", "Клиенты", metadata["unique_clients"], "kpi_value"),
        ("E6:G6", "E7:G8", "Строк клиент-месяц", metadata["panel_rows"], "kpi_value"),
        ("H6:J6", "H7:J8", "Штрафы 2025", metadata["fines_2025_total"], "kpi_value"),
        ("K6:M6", "K7:M8", "Штрафы 2026", metadata["fines_2026_total"], "kpi_value"),
        ("N6:Q6", "N7:Q8", "Операции топлива", metadata["fuel_transactions"], "kpi_value"),
    ]
    for label_range, value_range, label, value, value_format in kpis:
        overview.merge_range(label_range, label, fmt["kpi_label"])
        overview.merge_range(value_range, value, fmt[value_format])
    overview.merge_range("B10:Q11", f"Контроль качества: {metadata['fuel_negative_transactions']:,} операций с отрицательным объёмом; {metadata['fuel_low_price_transactions']:,} положительных операций дешевле 50 ₽/л; {metadata['fuel_high_price_transactions']:,} положительных операций дороже 130 ₽/л. Максимальная цена в данных — {metadata['price_max']:.0f} ₽/л; цены 500 ₽/л нет.", fmt["note"])

    monthly_cols = {name: idx for idx, name in enumerate(monthly.columns)}
    first, last = 1, len(monthly)
    add_chart(
        workbook, overview, "line", "Цена топлива: raw и робастный сценарий",
        None,
        [
            {"name": "Raw медиана", "categories": ["Месяцы", first, monthly_cols["month_name"], last, monthly_cols["month_name"]], "values": ["Месяцы", first, monthly_cols["fuel_price_median_raw"], last, monthly_cols["fuel_price_median_raw"]], "line": {"color": "#94A3B8", "width": 2.0}, "marker": {"type": "circle", "size": 5}},
            {"name": "50–130 ₽/л, медиана", "categories": ["Месяцы", first, monthly_cols["month_name"], last, monthly_cols["month_name"]], "values": ["Месяцы", first, monthly_cols["fuel_price_median_standard_band"], last, monthly_cols["fuel_price_median_standard_band"]], "line": {"color": "#0F766E", "width": 2.5}, "marker": {"type": "circle", "size": 6}},
        ],
        "B13", "₽/л"
    )
    add_chart(
        workbook, overview, "line", "Штрафы на клиента: одинаковые месяцы 2025 и 2026",
        None,
        [
            {"name": "2025", "categories": ["Месяцы", first, monthly_cols["month_name"], last, monthly_cols["month_name"]], "values": ["Месяцы", first, monthly_cols["fines_per_client_2025"], last, monthly_cols["fines_per_client_2025"]], "line": {"color": "#64748B", "width": 2.0}, "marker": {"type": "circle", "size": 5}},
            {"name": "2026", "categories": ["Месяцы", first, monthly_cols["month_name"], last, monthly_cols["month_name"]], "values": ["Месяцы", first, monthly_cols["fines_per_client_2026"], last, monthly_cols["fines_per_client_2026"]], "line": {"color": "#F59E0B", "width": 2.5}, "marker": {"type": "circle", "size": 6}},
        ],
        "J13", "штрафов / клиент"
    )
    add_chart(
        workbook, overview, "column", "Объём: все положительные покупки и робастный диапазон",
        None,
        [
            {"name": "Положительный объём", "categories": ["Месяцы", first, monthly_cols["month_name"], last, monthly_cols["month_name"]], "values": ["Месяцы", first, monthly_cols["fuel_volume_positive_l"], last, monthly_cols["fuel_volume_positive_l"]], "fill": {"color": "#CBD5E1"}, "border": {"none": True}},
            {"name": "50–130 ₽/л", "categories": ["Месяцы", first, monthly_cols["month_name"], last, monthly_cols["month_name"]], "values": ["Месяцы", first, monthly_cols["fuel_volume_standard_band_l"], last, monthly_cols["fuel_volume_standard_band_l"]], "fill": {"color": "#14B8A6"}, "border": {"none": True}},
        ],
        "B31", "литры"
    )
    add_chart(
        workbook, overview, "column", "Флаги качества топлива по месяцам",
        None,
        [
            {"name": "Отрицательный объём", "categories": ["Месяцы", first, monthly_cols["month_name"], last, monthly_cols["month_name"]], "values": ["Месяцы", first, monthly_cols["fuel_txn_negative"], last, monthly_cols["fuel_txn_negative"]], "fill": {"color": "#64748B"}, "border": {"none": True}},
            {"name": "Цена < 50", "categories": ["Месяцы", first, monthly_cols["month_name"], last, monthly_cols["month_name"]], "values": ["Месяцы", first, monthly_cols["fuel_txn_low_price"], last, monthly_cols["fuel_txn_low_price"]], "fill": {"color": "#FBBF24"}, "border": {"none": True}},
            {"name": "Цена > 130", "categories": ["Месяцы", first, monthly_cols["month_name"], last, monthly_cols["month_name"]], "values": ["Месяцы", first, monthly_cols["fuel_txn_high_price"], last, monthly_cols["fuel_txn_high_price"]], "fill": {"color": "#EF4444"}, "border": {"none": True}},
        ],
        "J31", "операции"
    )
    overview.merge_range("B49:Q51", "Важно: высокая цена может быть другим видом топлива или продуктом, а отрицательный объём — возвратом/корректировкой. Поэтому такие записи не объявлены ошибками: они сохранены в raw-полях и исключены только из отдельного робастного сценария.", fmt["note"])
    overview.freeze_panes(4, 1)
    overview.set_landscape()
    overview.fit_to_pages(1, 1)
    overview.set_margins(left=0.25, right=0.25, top=0.35, bottom=0.35)
    overview.print_area("B1:Q51")

    # Supporting tables.
    write_frame(monthly_sheet, monthly, fmt)
    monthly_sheet.set_column(0, len(monthly.columns) - 1, 15)
    monthly_sheet.set_column(monthly_cols["month_name"], monthly_cols["month_name"], 13)
    write_frame(region_sheet, regions, fmt)
    region_sheet.set_column(0, 0, 25)
    region_sheet.set_column(1, len(regions.columns) - 1, 16)
    write_frame(quality_sheet, quality, fmt)
    quality_sheet.set_column("A:A", 43)
    quality_sheet.set_column("B:B", 22)
    quality_sheet.set_column("C:C", 55)
    quality_sheet.write(len(quality) + 3, 0, "Распределения", fmt["section"])
    write_frame(quality_sheet, histograms, fmt, start_row=len(quality) + 5, autofilter=False)
    quality_sheet.set_column("D:D", 13)

    # Correlation heatmaps.
    for target, frame, title, note in [
        (corr_sheet, corr_level, "Матрица Спирмена: уровни client_id × месяц", "Парные корреляции, минимум 100 наблюдений. Это описательные связи, не причинные эффекты."),
        (change_corr_sheet, corr_change, "Матрица Спирмена: изменения до/во время кризиса", "Изменения сравнивают среднее за апрель–май со средним за июнь–август на уровне клиента."),
    ]:
        target.merge_range(0, 0, 0, len(frame.columns), title, fmt["title"])
        target.merge_range(1, 0, 2, len(frame.columns), note, fmt["subtitle"])
        target.write(4, 0, "Показатель", fmt["matrix_header"])
        for j, col in enumerate(frame.columns, start=1):
            target.write(4, j, col, fmt["matrix_header"])
        for i, (idx, row) in enumerate(frame.iterrows(), start=5):
            target.write(i, 0, idx, fmt["matrix_header"])
            for j, value in enumerate(row, start=1):
                target.write_number(i, j, float(value), fmt["matrix"])
        target.conditional_format(5, 1, 4 + len(frame), len(frame.columns), {
            "type": "3_color_scale", "min_type": "num", "min_value": -1, "min_color": "#2563EB",
            "mid_type": "num", "mid_value": 0, "mid_color": "#F8FAFC",
            "max_type": "num", "max_value": 1, "max_color": "#DC2626",
        })
        target.set_column(0, 0, 34)
        target.set_column(1, len(frame.columns), 13)
        target.set_row(4, 90)
        target.freeze_panes(5, 1)

    # Regional maps are static PNGs because native Excel charts do not support Russian GeoJSON choropleths.
    maps_sheet.set_column("A:A", 3)
    maps_sheet.set_column("B:Q", 12)
    maps_sheet.merge_range("B2:Q2", "Карты регионального распределения", fmt["title"])
    maps_sheet.merge_range("B3:Q4", "Регион определяется по регистрации клиента (КЛАДР), а не по координатам АЗС: в транзакциях нет местоположения станции. Поэтому карты показывают профиль клиентской базы.", fmt["subtitle"])
    map_paths = [
        FIG / "map_fines_per_client_2026.png",
        FIG / "map_fuel_price_standard.png",
        FIG / "map_fines_change_2026_vs_2025.png",
    ]
    for row, image_path in zip([5, 47, 89], map_paths):
        maps_sheet.insert_image(row, 1, str(image_path), {"x_scale": 0.65, "y_scale": 0.65})
    maps_sheet.freeze_panes(4, 1)
    maps_sheet.set_landscape()
    maps_sheet.fit_to_pages(1, 3)
    maps_sheet.set_margins(left=0.2, right=0.2, top=0.35, bottom=0.35)
    maps_sheet.print_area("B1:Q130")

    # Methodology and source links.
    methods_sheet.set_column("A:A", 31)
    methods_sheet.set_column("B:B", 80)
    methods_sheet.set_column("C:C", 34)
    methods_sheet.merge_range("A1:C1", "Методика и ограничения", fmt["title"])
    methods = [
        ("Единица наблюдения", "Одна строка = один client_id × календарный месяц, апрель–август 2026. Полная сетка включает месяцы без событий.", "128 375 строк"),
        ("Сопоставление", "Демография, штрафы и топливо соединены по client_id. Штрафы 2025 сравниваются с тем же месяцем 2026.", "Ключ: client_id + месяц"),
        ("Дубликаты автомобилей", "Сначала одна запись на client_id × auto_document_id; повторяющаяся история штрафов 2025 берётся max, затем суммируется по автомобилям клиента.", "Защита от двойного счёта"),
        ("Отрицательный объём", "Не удаляется. Хранится net-объём и отдельный абсолютный объём отрицательных операций. В метриках покупок используются только операции volume > 0.", "Флаг вероятного возврата/корректировки"),
        ("Цена топлива", "Raw-цены сохранены. Основной робастный сценарий: volume > 0 и 50 ≤ price ≤ 130 ₽/л. Значения вне диапазона не исправляются и доступны отдельными счётчиками.", "Сценарий чувствительности"),
        ("Экстремальный объём", f"Порог — 99,5-й перцентиль положительных операций: {metadata['extreme_volume_threshold']:.2f} л. Записи только помечаются.", "Не удалено"),
        ("Почему не удалять 20/359 ₽", "Без поля типа топлива нельзя доказать ошибку: дорогой кластер может быть другим продуктом, дешёвые значения — скидкой/субсидией/технической ценой. Проверяем выводы на raw и робастном сценариях.", "Максимум 359 ₽/л; 500 ₽/л отсутствует"),
        ("Региональные карты", "Регион — место регистрации клиента, не место заправки. Показатели топлива на карте описывают клиентов региона.", "16 регионов в выборке"),
        ("Корреляции", "Спирмен выбран из-за счётных показателей, асимметрии и выбросов. Корреляция не доказывает причинность; для вывода нужны модели с контролями и устойчивостью.", "Две матрицы: уровни и изменения"),
    ]
    for j, head in enumerate(["Раздел", "Решение", "Комментарий"]):
        methods_sheet.write(2, j, head, fmt["method_head"])
    for i, row in enumerate(methods, start=3):
        for j, value in enumerate(row):
            methods_sheet.write(i, j, value, fmt["text_wrap"])
        methods_sheet.set_row(i, 62)
    source_row = len(methods) + 5
    methods_sheet.write(source_row, 0, "Источники", fmt["section"])
    methods_sheet.write_url(source_row + 2, 0, "https://www.rosstat.gov.ru/storage/mediabank/128_26-08-2026.html", string="Росстат: цены, июль 2026")
    methods_sheet.write(source_row + 2, 1, "Средняя цена бензина — 77,14 ₽/л; АИ-92 — 73,21; АИ-95 — 79,44; АИ-98+ — 100,99; дизель — 92,51.", fmt["text_wrap"])
    methods_sheet.write_url(source_row + 3, 0, "https://github.com/rnekrasov-msk/geojson", string="Границы субъектов РФ")
    methods_sheet.write(source_row + 3, 1, "GeoJSON субъектов России; карта используется только как географическая подложка.", fmt["text_wrap"])

    # Large client-month table (streamed).
    if panel_sheet is not None:
        panel_sheet.set_tab_color("#0F766E")
        panel_sheet.freeze_panes(1, 2)
        panel_sheet.set_default_row(15)
        panel_path = WORK / "client_month_panel_quality.csv"
        text_cols = {"client_id", "month_name", "period", "gender", "age_type_code", "kladr_code", "registration_region", "fuel_quality_flag"}
        date_cols = {"month", "subscription_creation_date"}
        percent_cols = {"fines_change_pct_2026_vs_2025", "fuel_low_price_share_positive", "fuel_high_price_share_positive", "fuel_negative_txn_share_all"}
        formula_cols = {"fines_change_abs_2026_vs_2025", "fines_change_pct_2026_vs_2025", "fuel_price_weighted_standard_band", "fuel_low_price_share_positive", "fuel_high_price_share_positive", "fuel_negative_txn_share_all"}
        with panel_path.open("r", encoding="utf-8", newline="") as stream:
            reader = csv.reader(stream)
            headers = next(reader)
            positions = {name: i for i, name in enumerate(headers)}
            for col, header in enumerate(headers):
                if header.startswith("fuel_"):
                    header_format = fmt["header_fuel"]
                elif header.startswith("fine") or header.startswith("offence") or header.startswith("severe") or header.startswith("speeding"):
                    header_format = fmt["header_fines"]
                else:
                    header_format = fmt["header_demo"]
                panel_sheet.write(0, col, header, header_format)
                width = 29 if header == "client_id" else 24 if header == "registration_region" else 18 if header in date_cols else 16
                panel_sheet.set_column(col, col, width)
            panel_sheet.set_row(0, 68)
            last_row = 0
            for excel_row, row in enumerate(reader, start=1):
                last_row = excel_row
                for col, header in enumerate(headers):
                    raw = row[col] if col < len(row) else ""
                    if raw == "":
                        panel_sheet.write_blank(excel_row, col, None, fmt["text"])
                        continue
                    if header in formula_cols:
                        # Formulas use one-based Excel rows and cached values from the reproducible CSV pipeline.
                        r = excel_row + 1
                        if header == "fines_change_abs_2026_vs_2025":
                            formula = f"={get_column_letter(positions['fines_2026']+1)}{r}-{get_column_letter(positions['fines_2025']+1)}{r}"
                        elif header == "fines_change_pct_2026_vs_2025":
                            c26, c25 = get_column_letter(positions['fines_2026']+1), get_column_letter(positions['fines_2025']+1)
                            formula = f'=IFERROR(({c26}{r}-{c25}{r})/{c25}{r},"")'
                        elif header == "fuel_price_weighted_standard_band":
                            spend, volume = get_column_letter(positions['fuel_spend_standard_band_rub']+1), get_column_letter(positions['fuel_volume_standard_band_l']+1)
                            formula = f'=IFERROR({spend}{r}/{volume}{r},"")'
                        elif header == "fuel_low_price_share_positive":
                            num, den = get_column_letter(positions['fuel_txn_low_price']+1), get_column_letter(positions['fuel_txn_positive']+1)
                            formula = f'=IFERROR({num}{r}/{den}{r},"")'
                        elif header == "fuel_high_price_share_positive":
                            num, den = get_column_letter(positions['fuel_txn_high_price']+1), get_column_letter(positions['fuel_txn_positive']+1)
                            formula = f'=IFERROR({num}{r}/{den}{r},"")'
                        else:
                            num, den = get_column_letter(positions['fuel_txn_negative']+1), get_column_letter(positions['fuel_txn_all']+1)
                            formula = f'=IFERROR({num}{r}/{den}{r},"")'
                        value = float(raw)
                        panel_sheet.write_formula(excel_row, col, formula, fmt["percent"] if header in percent_cols else fmt["decimal"], value)
                    elif header in date_cols:
                        panel_sheet.write_datetime(excel_row, col, datetime.strptime(raw[:10], "%Y-%m-%d"), fmt["date"])
                    elif header in text_cols:
                        panel_sheet.write_string(excel_row, col, raw, fmt["text"])
                    else:
                        number = float(raw)
                        if header in percent_cols:
                            number_format = fmt["percent"]
                        elif any(token in header for token in ("price", "amount", "spend")):
                            number_format = fmt["money"]
                        elif any(token in header for token in ("volume", "engine_", "vehicle_year")):
                            number_format = fmt["decimal"]
                        else:
                            number_format = fmt["integer"]
                        panel_sheet.write_number(excel_row, col, number, number_format)
            panel_sheet.autofilter(0, 0, last_row, len(headers) - 1)
            for col_name, colors in [
                ("fines_change_pct_2026_vs_2025", ("#DBEAFE", "#FFFFFF", "#FECACA")),
                ("fuel_quality_flag", None),
            ]:
                col = positions[col_name]
                if colors:
                    panel_sheet.conditional_format(1, col, last_row, col, {"type": "3_color_scale", "min_color": colors[0], "mid_type": "num", "mid_value": 0, "mid_color": colors[1], "max_color": colors[2]})
                else:
                    panel_sheet.conditional_format(1, col, last_row, col, {"type": "text", "criteria": "containing", "value": "Есть флаги", "format": workbook.add_format({"bg_color": "#FEF3C7", "font_color": "#92400E"})})

    workbook.close()


def verify(path: Path, expected_panel: bool):
    book = load_workbook(path, read_only=False, data_only=False)
    expected = ["Обзор", "Месяцы", "Регионы", "Выбросы", "Корреляции", "Корр_изменений", "Карты", "Методика"]
    if expected_panel:
        expected.append("Клиент-месяц")
    assert book.sheetnames == expected
    assert book["Обзор"].freeze_panes == "B5"
    assert len(book["Обзор"]._charts) == 4
    assert len(book["Карты"]._images) == 3
    if expected_panel:
        ws = book["Клиент-месяц"]
        assert ws.max_row == 128376
        assert ws.max_column == 52
        assert ws.freeze_panes == "C2"
        assert ws.auto_filter.ref.endswith(str(ws.max_row))
        assert ws.cell(2, 1).value
        assert ws.cell(ws.max_row, 1).value
    book.close()


def main():
    shutil.copyfile(WORK / "client_month_panel_quality.csv", OUTPUT_CSV)
    build_workbook(OUTPUT, include_panel=True)
    verify(OUTPUT, expected_panel=True)
    build_workbook(PREVIEW, include_panel=False)
    verify(PREVIEW, expected_panel=False)
    metadata = json.loads((WORK / "metadata.json").read_text(encoding="utf-8"))
    print(json.dumps({"workbook": str(OUTPUT), "csv": str(OUTPUT_CSV), "preview": str(PREVIEW), **metadata}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
