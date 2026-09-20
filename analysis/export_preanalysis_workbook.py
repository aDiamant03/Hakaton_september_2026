from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
import xlsxwriter
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / "working_data" / "preanalysis"
FIG = ROOT / "figures" / "preanalysis"
OUT = ROOT / "outputs" / "preanalysis"
TMP = ROOT / "tmp" / "preanalysis"
OUT.mkdir(parents=True, exist_ok=True)
TMP.mkdir(parents=True, exist_ok=True)
OUTPUT = OUT / "DANO_преданализ_гипотезы.xlsx"
PREVIEW = TMP / "preview.xlsx"


def formats(workbook):
    return {
        "title": workbook.add_format({"font_name": "Arial", "font_size": 23, "bold": True, "font_color": "#111827"}),
        "subtitle": workbook.add_format({"font_name": "Arial", "font_size": 11, "font_color": "#475569", "text_wrap": True, "valign": "top"}),
        "section": workbook.add_format({"font_name": "Arial", "font_size": 14, "bold": True, "font_color": "#111827", "bg_color": "#FEF3C7", "bottom": 2, "bottom_color": "#F59E0B"}),
        "header": workbook.add_format({"font_name": "Arial", "font_size": 9, "bold": True, "font_color": "#FFFFFF", "bg_color": "#334155", "align": "center", "valign": "vcenter", "text_wrap": True, "border": 1, "border_color": "#FFFFFF"}),
        "text": workbook.add_format({"font_name": "Arial", "font_size": 9, "font_color": "#1F2937", "valign": "top"}),
        "wrap": workbook.add_format({"font_name": "Arial", "font_size": 9, "font_color": "#1F2937", "valign": "top", "text_wrap": True}),
        "integer": workbook.add_format({"font_name": "Arial", "font_size": 9, "num_format": "#,##0", "valign": "top"}),
        "decimal": workbook.add_format({"font_name": "Arial", "font_size": 9, "num_format": "0.000", "valign": "top"}),
        "number2": workbook.add_format({"font_name": "Arial", "font_size": 9, "num_format": "#,##0.00", "valign": "top"}),
        "percent": workbook.add_format({"font_name": "Arial", "font_size": 9, "num_format": "0.0%", "valign": "top"}),
        "pvalue": workbook.add_format({"font_name": "Arial", "font_size": 9, "num_format": "0.0000", "valign": "top"}),
        "kpi_label": workbook.add_format({"font_name": "Arial", "font_size": 10, "font_color": "#475569", "align": "center", "valign": "vcenter", "border": 1, "border_color": "#E2E8F0", "bg_color": "#FFFFFF"}),
        "kpi_value": workbook.add_format({"font_name": "Arial", "font_size": 18, "bold": True, "font_color": "#0F766E", "align": "center", "valign": "vcenter", "border": 1, "border_color": "#E2E8F0", "bg_color": "#FFFFFF", "num_format": "0.0%"}),
        "kpi_did": workbook.add_format({"font_name": "Arial", "font_size": 18, "bold": True, "font_color": "#2563EB", "align": "center", "valign": "vcenter", "border": 1, "border_color": "#E2E8F0", "bg_color": "#FFFFFF", "num_format": "0.000"}),
        "note": workbook.add_format({"font_name": "Arial", "font_size": 10, "font_color": "#475569", "text_wrap": True, "valign": "top", "border": 1, "border_color": "#E2E8F0", "bg_color": "#F8FAFC"}),
        "good": workbook.add_format({"font_name": "Arial", "font_size": 9, "font_color": "#166534", "bg_color": "#DCFCE7", "bold": True}),
        "warn": workbook.add_format({"font_name": "Arial", "font_size": 9, "font_color": "#92400E", "bg_color": "#FEF3C7", "bold": True}),
        "bad": workbook.add_format({"font_name": "Arial", "font_size": 9, "font_color": "#991B1B", "bg_color": "#FEE2E2", "bold": True}),
    }


def setup(sheet, zoom=90):
    sheet.hide_gridlines(2)
    sheet.set_zoom(zoom)
    sheet.set_default_row(18)


def write_frame(sheet, frame, fmt, start_row=0, start_col=0, filter_on=True):
    for j, header in enumerate(frame.columns):
        sheet.write(start_row, start_col + j, header, fmt["header"])
    sheet.set_row(start_row, 44)
    for i, row in enumerate(frame.itertuples(index=False, name=None), start=start_row + 1):
        for j, value in enumerate(row):
            column = frame.columns[j]
            if pd.isna(value):
                sheet.write_blank(i, start_col + j, None, fmt["text"])
            elif isinstance(value, bool):
                sheet.write(i, start_col + j, "Да" if value else "Нет", fmt["text"])
            elif isinstance(value, (int, float)):
                if "pct" in column or "share" in column:
                    number_format = fmt["percent"]
                elif "pvalue" in column:
                    number_format = fmt["pvalue"]
                elif isinstance(value, int) or column in {"clients", "orders", "months", "value"}:
                    number_format = fmt["integer"]
                else:
                    number_format = fmt["decimal"]
                sheet.write_number(i, start_col + j, float(value), number_format)
            else:
                sheet.write(i, start_col + j, str(value), fmt["wrap"])
    if filter_on and len(frame):
        sheet.autofilter(start_row, start_col, start_row + len(frame), start_col + len(frame.columns) - 1)
    sheet.freeze_panes(start_row + 1, start_col)


def add_chart_style(chart, title, y_name=""):
    chart.set_title({"name": title, "name_font": {"name": "Arial", "size": 12, "bold": True}})
    chart.set_x_axis({"num_font": {"name": "Arial", "size": 8}, "name_font": {"name": "Arial", "size": 9}})
    chart.set_y_axis({"name": y_name, "major_gridlines": {"visible": True, "line": {"color": "#E2E8F0"}}, "num_font": {"name": "Arial", "size": 8}, "name_font": {"name": "Arial", "size": 9}})
    chart.set_legend({"position": "bottom", "font": {"name": "Arial", "size": 8}})
    chart.set_plotarea({"fill": {"color": "#FFFFFF"}, "border": {"none": True}})
    chart.set_chartarea({"fill": {"color": "#FFFFFF"}, "border": {"color": "#E2E8F0"}})


def build(path: Path, include_clients=True):
    findings = json.loads((WORK / "findings.json").read_text(encoding="utf-8"))
    event = pd.read_csv(WORK / "monthly_event_study.csv")
    fuel = pd.read_csv(WORK / "monthly_fuel_indices.csv")
    period_fuel = pd.read_csv(WORK / "period_fuel_summary.csv")
    regions = pd.read_csv(WORK / "region_hypothesis.csv")
    segments = pd.read_csv(WORK / "fuel_intensity_segments.csv")
    offence = pd.read_csv(WORK / "offence_composition.csv")
    quality = pd.read_csv(WORK / "data_quality_preanalysis.csv")
    validation = pd.read_csv(WORK / "map_validation.csv", dtype={"kladr_code": "string"})
    hypotheses = pd.read_csv(WORK / "hypotheses_preanalysis.csv")
    exposure_segments = pd.read_csv(WORK / "fuel_exposure_segments.csv")
    exposure_robustness = pd.read_csv(WORK / "fuel_exposure_robustness.csv")
    vehicle_segments = pd.read_csv(WORK / "vehicle_extreme_segments.csv")
    new_variables = pd.read_csv(WORK / "new_variables_dictionary.csv")
    clients = pd.read_csv(WORK / "client_did.csv") if include_clients else None

    pre = period_fuel.set_index("period").loc["До кризиса"]
    crisis = period_fuel.set_index("period").loc["Кризис"]
    price_change = crisis["price_standard_weighted"] / pre["price_standard_weighted"] - 1
    volume_change = crisis["volume_per_client_month_l"] / pre["volume_per_client_month_l"] - 1
    orders_change = crisis["orders_per_client_month"] / pre["orders_per_client_month"] - 1

    workbook = xlsxwriter.Workbook(path, {"constant_memory": include_clients, "strings_to_numbers": False, "nan_inf_to_errors": True})
    workbook.set_properties({"title": "DANO — преданализ топливного кризиса", "subject": "Гипотезы, исследовательский дизайн и проверка карт", "company": "DATA SWAGA team"})
    fmt = formats(workbook)
    overview = workbook.add_worksheet("Преданализ")
    favorite_sheet = workbook.add_worksheet("Фаворит")
    hypothesis_sheet = workbook.add_worksheet("Гипотезы")
    variables_sheet = workbook.add_worksheet("Новые переменные")
    dynamics = workbook.add_worksheet("Динамика")
    region_sheet = workbook.add_worksheet("Регионы")
    maps_sheet = workbook.add_worksheet("Карты")
    quality_sheet = workbook.add_worksheet("Качество")
    model_sheet = workbook.add_worksheet("План модели")
    client_sheet = workbook.add_worksheet("Клиентский DiD") if include_clients else None
    for sheet in [overview, favorite_sheet, hypothesis_sheet, variables_sheet, dynamics, region_sheet, maps_sheet, quality_sheet, model_sheet] + ([client_sheet] if client_sheet else []):
        setup(sheet, 85 if sheet is client_sheet else 90)

    # Main pre-analysis view.
    overview.set_column("A:A", 3)
    overview.set_column("B:Q", 12)
    overview.set_row(1, 36)
    overview.merge_range("B2:Q2", "DANO — преданализ: что проверяем и что уже видно", fmt["title"])
    overview.merge_range("B3:Q4", "Главный вопрос нельзя проверять простым сравнением уровней 2025 и 2026. Основная метрика — разность разностей: насколько разрыв 2026–2025 изменился после мая. Топливные показатели рассматриваются и как индикатор кризиса, и как возможный механизм через изменение мобильности.", fmt["subtitle"])
    cards = [
        ("B6:E6", "B7:E8", "Fine DiD, штрафов/клиент/месяц", findings["fine_did_mean"], "kpi_did"),
        ("F6:I6", "F7:I8", "Робастная цена: кризис к pre", price_change, "kpi_value"),
        ("J6:M6", "J7:M8", "Объём/клиент/месяц", volume_change, "kpi_value"),
        ("N6:Q6", "N7:Q8", "Заказы/клиент/месяц", orders_change, "kpi_value"),
    ]
    for label_range, value_range, label, value, style in cards:
        overview.merge_range(label_range, label, fmt["kpi_label"])
        overview.merge_range(value_range, value, fmt[style])
    overview.merge_range("B10:Q11", f"Предварительный общий эффект: {findings['fine_did_mean']:+.3f} штрафа на клиента в месяц, 95% ДИ [{findings['fine_did_ci95_low']:+.3f}; {findings['fine_did_ci95_high']:+.3f}]. Это не окончательная причинная оценка: ещё нужны проверки трендов, спецификаций и региональной устойчивости.", fmt["note"])

    # Event-study chart.
    event_cols = {name: index for index, name in enumerate(event.columns)}
    event_chart = workbook.add_chart({"type": "line"})
    event_chart.add_series({
        "name": "2026 минус 2025",
        "categories": ["Динамика", 1, event_cols["month_name"], len(event), event_cols["month_name"]],
        "values": ["Динамика", 1, event_cols["yoy_delta_per_client"], len(event), event_cols["yoy_delta_per_client"]],
        "line": {"color": "#2563EB", "width": 2.5}, "marker": {"type": "circle", "size": 7, "border": {"color": "#2563EB"}, "fill": {"color": "#FFFFFF"}},
        "y_error_bars": {"type": "custom", "plus_values": f"=Динамика!${xlsxwriter.utility.xl_col_to_name(event_cols['ci95_plus'])}$2:${xlsxwriter.utility.xl_col_to_name(event_cols['ci95_plus'])}$6", "minus_values": f"=Динамика!${xlsxwriter.utility.xl_col_to_name(event_cols['ci95_minus'])}$2:${xlsxwriter.utility.xl_col_to_name(event_cols['ci95_minus'])}$6"},
    })
    add_chart_style(event_chart, "Разрыв штрафов 2026–2025 по месяцам", "штрафов / клиент")
    event_chart.set_y_axis({"name": "штрафов / клиент", "min": -0.7, "max": 0.05, "major_unit": 0.1, "num_format": "0.0", "crossing": 0, "major_gridlines": {"visible": True, "line": {"color": "#E2E8F0"}}})
    overview.insert_chart("B13", event_chart, {"x_scale": 1.22, "y_scale": 1.2})

    # Fuel indices chart.
    fuel_cols = {name: index for index, name in enumerate(fuel.columns)}
    fuel_chart = workbook.add_chart({"type": "line"})
    for name, column, color in [
        ("Цена", "price_standard_weighted_index_april", "#F59E0B"),
        ("Объём/клиент", "volume_per_client_l_index_april", "#0F766E"),
        ("Заказы/клиент", "orders_per_client_index_april", "#2563EB"),
        ("Литров/заказ", "liters_per_order_index_april", "#7C3AED"),
    ]:
        fuel_chart.add_series({"name": name, "categories": ["Динамика", 9, fuel_cols["month_name"], 13, fuel_cols["month_name"]], "values": ["Динамика", 9, fuel_cols[column], 13, fuel_cols[column]], "line": {"color": color, "width": 2.2}, "marker": {"type": "circle", "size": 5}})
    add_chart_style(fuel_chart, "Цена и топливная активность (апрель = 100)", "индекс")
    overview.insert_chart("J13", fuel_chart, {"x_scale": 1.22, "y_scale": 1.2})

    # Regional scatter.
    region_cols = {name: index for index, name in enumerate(regions.columns)}
    scatter = workbook.add_chart({"type": "scatter", "subtype": "straight_with_markers"})
    scatter.add_series({
        "name": "16 регионов",
        "categories": ["Регионы", 1, region_cols["price_change_pct"], len(regions), region_cols["price_change_pct"]],
        "values": ["Регионы", 1, region_cols["fine_did_per_client_month"], len(regions), region_cols["fine_did_per_client_month"]],
        "line": {"none": True}, "marker": {"type": "circle", "size": 7, "fill": {"color": "#14B8A6", "transparency": 20}, "border": {"color": "#0F766E"}},
        "trendline": {"type": "linear", "line": {"color": "#64748B", "dash_type": "dash"}},
    })
    add_chart_style(scatter, "Региональный рост цены и fine DiD", "fine DiD")
    scatter.set_x_axis({"name": "рост цены после мая", "num_format": "0%", "major_gridlines": {"visible": True, "line": {"color": "#E2E8F0"}}})
    overview.insert_chart("B31", scatter, {"x_scale": 1.22, "y_scale": 1.2})

    # Heterogeneity by baseline fuel intensity.
    segment_cols = {name: index for index, name in enumerate(segments.columns)}
    segment_chart = workbook.add_chart({"type": "column"})
    segment_chart.add_series({
        "name": "Fine DiD",
        "categories": ["Динамика", 18, segment_cols["segment"], 22, segment_cols["segment"]],
        "values": ["Динамика", 18, segment_cols["fine_did_mean"], 22, segment_cols["fine_did_mean"]],
        "fill": {"color": "#2563EB"}, "border": {"none": True},
        "y_error_bars": {"type": "custom", "plus_values": f"=Динамика!${xlsxwriter.utility.xl_col_to_name(segment_cols['ci95_plus'])}$19:${xlsxwriter.utility.xl_col_to_name(segment_cols['ci95_plus'])}$23", "minus_values": f"=Динамика!${xlsxwriter.utility.xl_col_to_name(segment_cols['ci95_minus'])}$19:${xlsxwriter.utility.xl_col_to_name(segment_cols['ci95_minus'])}$23"},
    })
    add_chart_style(segment_chart, "Эффект по докризисной топливной активности", "fine DiD")
    segment_chart.set_legend({"none": True})
    overview.insert_chart("J31", segment_chart, {"x_scale": 1.22, "y_scale": 1.2})
    overview.merge_range("B49:Q51", f"Пока нет подтверждения, что больший региональный рост цены связан с более сильным снижением штрафов: Spearman ρ={findings['region_price_fines_spearman']:.2f}, p={findings['region_price_fines_pvalue']:.3f}. На уровне клиента связь изменения объёма и fine DiD статистически заметна, но очень мала: ρ={findings['client_volume_fine_did_spearman']:.3f}. Поэтому главный следующий шаг — модель, а не новая корреляционная матрица.", fmt["note"])
    overview.freeze_panes(4, 1)
    overview.set_landscape()
    overview.fit_to_pages(1, 1)
    overview.set_margins(left=0.25, right=0.25, top=0.35, bottom=0.35)
    overview.print_area("B1:Q51")

    # Winning hypothesis: pre-crisis fuel exposure relative to vehicle value.
    favorite_sheet.set_column("A:A", 3)
    favorite_sheet.set_column("B:Q", 12)
    favorite_sheet.set_row(1, 36)
    favorite_sheet.merge_range("B2:Q2", "Фаворит: топливная нагрузка и сокращение мобильности", fmt["title"])
    favorite_sheet.merge_range(
        "B3:Q5",
        "Гипотеза: после начала кризиса штрафов стало меньше через сокращение мобильности, причём эффект сильнее у клиентов с высокой докризисной топливной нагрузкой. Нагрузка определяется только по докризисным данным: средние расходы на обычное топливо в апреле–мае, делённые на медианную стоимость автомобиля клиента. Это прокси интенсивности и ценовой чувствительности, а не прямое измерение дохода.",
        fmt["subtitle"],
    )
    favorite_cards = [
        ("B7:E7", "B8:E9", "Fine DiD: Q1 нагрузки", exposure_segments.iloc[0]["fine_did_mean"], "kpi_did"),
        ("F7:I7", "F8:I9", "Fine DiD: Q4 нагрузки", exposure_segments.iloc[-1]["fine_did_mean"], "kpi_did"),
        ("J7:M7", "J8:M9", "Разница Q4 − Q1", findings["fuel_exposure_q4_minus_q1"], "kpi_did"),
        ("N7:Q7", "N8:Q9", "p-value Q4 vs Q1", findings["fuel_exposure_q4_vs_q1_pvalue"], "kpi_did"),
    ]
    for label_range, value_range, label, value, style in favorite_cards:
        favorite_sheet.merge_range(label_range, label, fmt["kpi_label"])
        favorite_sheet.merge_range(value_range, value, fmt[style])

    exposure_start = 11
    vehicle_start = exposure_start + len(exposure_segments) + 3
    write_frame(favorite_sheet, exposure_segments, fmt, start_row=exposure_start, start_col=1, filter_on=False)
    write_frame(favorite_sheet, vehicle_segments, fmt, start_row=vehicle_start, start_col=1, filter_on=False)
    exp_cols = {name: 1 + index for index, name in enumerate(exposure_segments.columns)}
    veh_cols = {name: 1 + index for index, name in enumerate(vehicle_segments.columns)}
    exp_chart = workbook.add_chart({"type": "column"})
    exp_chart.add_series({
        "name": "Fine DiD",
        "categories": ["Фаворит", exposure_start + 1, exp_cols["fuel_exposure_quartile"], exposure_start + len(exposure_segments), exp_cols["fuel_exposure_quartile"]],
        "values": ["Фаворит", exposure_start + 1, exp_cols["fine_did_mean"], exposure_start + len(exposure_segments), exp_cols["fine_did_mean"]],
        "fill": {"color": "#2563EB"}, "border": {"none": True},
        "y_error_bars": {"type": "custom", "plus_values": f"=Фаворит!${xlsxwriter.utility.xl_col_to_name(exp_cols['ci95_plus'])}${exposure_start+2}:${xlsxwriter.utility.xl_col_to_name(exp_cols['ci95_plus'])}${exposure_start+1+len(exposure_segments)}", "minus_values": f"=Фаворит!${xlsxwriter.utility.xl_col_to_name(exp_cols['ci95_minus'])}${exposure_start+2}:${xlsxwriter.utility.xl_col_to_name(exp_cols['ci95_minus'])}${exposure_start+1+len(exposure_segments)}"},
        "data_labels": {"value": True, "num_format": "0.000"},
    })
    add_chart_style(exp_chart, "Fine DiD по квартилям топливной нагрузки", "fine DiD")
    exp_chart.set_y_axis({"name": "fine DiD", "min": -0.25, "max": 0.02, "major_unit": 0.04, "num_format": "0.00", "major_gridlines": {"visible": True, "line": {"color": "#E2E8F0"}}})
    exp_chart.set_legend({"none": True})
    favorite_sheet.insert_chart("B25", exp_chart, {"x_scale": 1.22, "y_scale": 1.2})

    vehicle_chart = workbook.add_chart({"type": "column"})
    vehicle_chart.add_series({
        "name": "Fine DiD",
        "categories": ["Фаворит", vehicle_start + 1, veh_cols["vehicle_segment"], vehicle_start + len(vehicle_segments), veh_cols["vehicle_segment"]],
        "values": ["Фаворит", vehicle_start + 1, veh_cols["fine_did_mean"], vehicle_start + len(vehicle_segments), veh_cols["fine_did_mean"]],
        "fill": {"color": "#94A3B8"}, "border": {"none": True},
        "y_error_bars": {"type": "custom", "plus_values": f"=Фаворит!${xlsxwriter.utility.xl_col_to_name(veh_cols['ci95_plus'])}${vehicle_start+2}:${xlsxwriter.utility.xl_col_to_name(veh_cols['ci95_plus'])}${vehicle_start+1+len(vehicle_segments)}", "minus_values": f"=Фаворит!${xlsxwriter.utility.xl_col_to_name(veh_cols['ci95_minus'])}${vehicle_start+2}:${xlsxwriter.utility.xl_col_to_name(veh_cols['ci95_minus'])}${vehicle_start+1+len(vehicle_segments)}"},
        "data_labels": {"value": True, "num_format": "0.000"},
    })
    add_chart_style(vehicle_chart, "Простое сравнение автомобилей", "fine DiD")
    vehicle_chart.set_y_axis({"name": "fine DiD", "min": -0.18, "max": 0.02, "major_unit": 0.02, "num_format": "0.00", "major_gridlines": {"visible": True, "line": {"color": "#E2E8F0"}}})
    vehicle_chart.set_legend({"none": True})
    favorite_sheet.insert_chart("J25", vehicle_chart, {"x_scale": 1.22, "y_scale": 1.2})

    favorite_sheet.merge_range(
        "B43:Q45",
        f"Почему фаворит сильнее идеи old/cheap vs new/expensive: эффект по нагрузке монотонный и сохраняется в проверках. Q4−Q1 = {findings['fuel_exposure_q4_minus_q1']:+.3f}, p={findings['fuel_exposure_q4_vs_q1_pvalue']:.4f}. Простая разница старых дешёвых и новых дорогих машин мала: {findings['old_cheap_fine_did']:+.3f} против {findings['new_expensive_fine_did']:+.3f}, p={findings['old_cheap_vs_new_expensive_pvalue']:.3f}. Значит, автомобиль полезен как часть сегментации, но не как единственное объяснение.",
        fmt["note"],
    )
    favorite_sheet.write(47, 1, "Проверки устойчивости", fmt["section"])
    write_frame(favorite_sheet, exposure_robustness, fmt, start_row=49, start_col=1, filter_on=False)
    favorite_sheet.set_column("B:B", 31)
    favorite_sheet.set_column("C:I", 17)
    favorite_sheet.freeze_panes(5, 1)
    favorite_sheet.set_landscape()
    favorite_sheet.fit_to_pages(1, 1)
    favorite_sheet.set_margins(left=0.25, right=0.25, top=0.35, bottom=0.35)
    favorite_sheet.print_area("B1:Q54")

    write_frame(variables_sheet, new_variables, fmt)
    variables_sheet.set_column("A:A", 39)
    variables_sheet.set_column("B:B", 58)
    variables_sheet.set_column("C:C", 25)
    variables_sheet.set_column("D:D", 58)
    variables_sheet.set_default_row(52)

    # Hypotheses table.
    write_frame(hypothesis_sheet, hypotheses, fmt)
    hypothesis_sheet.set_column("A:A", 8)
    hypothesis_sheet.set_column("B:C", 52)
    hypothesis_sheet.set_column("D:D", 55)
    hypothesis_sheet.set_column("E:E", 26)
    hypothesis_sheet.set_default_row(72)
    status_col = hypotheses.columns.get_loc("preliminary_status")
    hypothesis_sheet.conditional_format(1, status_col, len(hypotheses), status_col, {"type": "text", "criteria": "containing", "value": "Поддерживается", "format": fmt["good"]})
    hypothesis_sheet.conditional_format(1, status_col, len(hypotheses), status_col, {"type": "text", "criteria": "containing", "value": "слабая", "format": fmt["warn"]})
    hypothesis_sheet.conditional_format(1, status_col, len(hypotheses), status_col, {"type": "text", "criteria": "containing", "value": "не поддерживается", "format": fmt["bad"]})

    # Supporting data on one sheet at known row locations used by charts.
    write_frame(dynamics, event, fmt, start_row=0)
    write_frame(dynamics, fuel, fmt, start_row=8, filter_on=False)
    write_frame(dynamics, segments, fmt, start_row=17, filter_on=False)
    write_frame(dynamics, offence, fmt, start_row=26, filter_on=False)
    write_frame(dynamics, period_fuel, fmt, start_row=33, filter_on=False)
    dynamics.set_column(0, 0, 23)
    dynamics.set_column(1, 30, 17)
    offence_cols = {name: index for index, name in enumerate(offence.columns)}
    offence_chart = workbook.add_chart({"type": "column"})
    for name, column, color in [
        ("Превышение скорости", "speeding_share", "#2563EB"),
        ("Тяжёлые нарушения", "severe_share", "#F59E0B"),
    ]:
        offence_chart.add_series({
            "name": name,
            "categories": ["Динамика", 27, offence_cols["period"], 28, offence_cols["period"]],
            "values": ["Динамика", 27, offence_cols[column], 28, offence_cols[column]],
            "fill": {"color": color}, "border": {"none": True},
            "data_labels": {"value": True, "num_format": "0.0%"},
        })
    add_chart_style(offence_chart, "Состав нарушений до и во время кризиса", "доля штрафов")
    offence_chart.set_y_axis({"name": "доля штрафов", "min": 0, "max": 1, "num_format": "0%", "major_gridlines": {"visible": True, "line": {"color": "#E2E8F0"}}})
    dynamics.insert_chart("M27", offence_chart, {"x_scale": 1.15, "y_scale": 1.1})

    write_frame(region_sheet, regions, fmt)
    region_sheet.set_column(0, 0, 27)
    region_sheet.set_column(1, len(regions.columns) - 1, 17)

    # Maps and exact polygon validation.
    maps_sheet.set_column("A:A", 3)
    maps_sheet.set_column("B:Q", 12)
    maps_sheet.set_row(1, 36)
    maps_sheet.merge_range("B2:Q2", "Карты, которые проверяют гипотезу", fmt["title"])
    maps_sheet.merge_range("B3:Q4", "Закрашены только 16 субъектов из КЛАДР. Для каждого названия проверено совпадение с полем NL_NAME_1 в GeoJSON и попадание контрольной точки регионального центра внутрь соответствующего полигона.", fmt["subtitle"])
    maps_sheet.insert_image(5, 1, str(FIG / "map_region_fine_did.png"), {"x_scale": 0.65, "y_scale": 0.65})
    maps_sheet.insert_image(47, 1, str(FIG / "map_region_price_change.png"), {"x_scale": 0.65, "y_scale": 0.65})
    validation_start = 90
    maps_sheet.write(validation_start, 1, "Проверка соответствия названий и полигонов", fmt["section"])
    write_frame(maps_sheet, validation, fmt, start_row=validation_start + 2, start_col=1, filter_on=False)
    maps_sheet.set_column("B:B", 12)
    maps_sheet.set_column("C:D", 32)
    maps_sheet.set_column("E:E", 20)
    maps_sheet.set_column("F:I", 17)
    status_idx = 1 + validation.columns.get_loc("validation_status")
    maps_sheet.conditional_format(validation_start + 3, status_idx, validation_start + 2 + len(validation), status_idx, {"type": "text", "criteria": "containing", "value": "OK", "format": fmt["good"]})
    maps_sheet.set_landscape()
    maps_sheet.fit_to_pages(1, 3)
    maps_sheet.freeze_panes(4, 1)

    write_frame(quality_sheet, quality, fmt)
    quality_sheet.set_column("A:A", 46)
    quality_sheet.set_column("B:B", 18)
    quality_sheet.set_column("C:C", 62)
    quality_sheet.set_default_row(40)

    # Concrete model plan.
    model_sheet.set_column("A:A", 28)
    model_sheet.set_column("B:B", 88)
    model_sheet.set_column("C:C", 40)
    model_sheet.merge_range("A1:C1", "План основного исследования", fmt["title"])
    model_rows = [
        ("Основной outcome", "Количество штрафов client_id × месяц. Нули сохраняются: отсутствие штрафа — наблюдаемый результат.", "fines_count"),
        ("Основное воздействие", "Post = июнь–август 2026. Сравнение с апрелем–маем и с теми же календарными месяцами 2025.", "post × year_2026"),
        ("Главная спецификация", "Y_it = α_i + λ_month + β·(year_2026 × post) + ε_it. β — изменение после мая сверх сезонного профиля 2025.", "Двухсторонние fixed effects / event-study"),
        ("Интенсивность кризиса", "Добавить региональный рост робастной цены и взаимодействие с post. Это проверяет H2, но цена может быть эндогенной.", "price_shock_region × post"),
        ("Контроли", "Эффекты клиента поглощают пол, возраст, регион и постоянные характеристики авто. В динамике добавить возраст авто/мощность × post, стаж подписки и региональные тренды.", "Не дублировать time-invariant признаки"),
        ("Механизм", "Объём и число заправок — возможные медиаторы. Не включать их в модель общего эффекта; оценить отдельно связь crisis → fuel и fuel → fines.", "Отдельная модель H3"),
        ("Оценивание", "Основной результат — линейная DiD для интерпретации. Устойчивость: Poisson PML/negative binomial для счётного outcome.", "Кластерные SE по клиенту"),
        ("Региональная неопределённость", "Всего 16 регионов: обычные кластерные SE по региону ненадёжны. Нужны leave-one-region-out и wild cluster bootstrap.", "Обязательная устойчивость"),
        ("Выбросы топлива", "Повторить модели для raw, положительных операций, диапазонов 50–130 и 40–150 ₽/л. Отрицательные объёмы учитывать отдельно.", "Sensitivity analysis"),
        ("Плацебо", "Считать ложной датой кризиса май и сравнить апрель→май. Сильный плацебо-эффект ослабит причинную интерпретацию.", "Проверка параллельных трендов"),
        ("Подвыборки", "Квартиль baseline fuel intensity, регион, возрастная группа, мощность автомобиля, клиенты с одной/несколькими машинами.", "H5 и устойчивость"),
    ]
    for j, head in enumerate(["Элемент", "Решение", "Роль"]):
        model_sheet.write(2, j, head, fmt["header"])
    for i, row in enumerate(model_rows, start=3):
        for j, value in enumerate(row):
            model_sheet.write(i, j, value, fmt["wrap"])
        model_sheet.set_row(i, 62)

    # Reproducible client-level DiD table.
    if client_sheet is not None:
        write_frame(client_sheet, clients, fmt)
        client_sheet.set_column(0, 0, 30)
        client_sheet.set_column(1, len(clients.columns) - 1, 19)

    workbook.close()


def verify(path: Path, include_clients: bool):
    book = load_workbook(path, read_only=False, data_only=False)
    expected = ["Преданализ", "Фаворит", "Гипотезы", "Новые переменные", "Динамика", "Регионы", "Карты", "Качество", "План модели"]
    if include_clients:
        expected.append("Клиентский DiD")
    assert book.sheetnames == expected
    assert len(book["Преданализ"]._charts) == 4
    assert len(book["Фаворит"]._charts) == 2
    assert len(book["Карты"]._images) == 2
    assert all(book["Карты"].cell(row, 10).value == "OK" for row in range(94, 110))
    if include_clients:
        assert book["Клиентский DiD"].max_row == 25676
    book.close()


def main():
    build(OUTPUT, include_clients=True)
    verify(OUTPUT, include_clients=True)
    build(PREVIEW, include_clients=False)
    verify(PREVIEW, include_clients=False)
    print(json.dumps({"output": str(OUTPUT), "preview": str(PREVIEW)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
