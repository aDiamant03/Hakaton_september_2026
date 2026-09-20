from __future__ import annotations

import csv
import json
from datetime import datetime
from pathlib import Path

import xlsxwriter
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "tmp" / "client_master" / "client_master.tsv"
SCHEMA_PATH = ROOT / "tmp" / "client_master" / "schema.json"
OUTPUT_DIR = ROOT / "outputs" / "client_master"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
OUTPUT = OUTPUT_DIR / "DANO_общая_таблица_по_клиентам.xlsx"

schema = json.loads(SCHEMA_PATH.read_text(encoding="utf-8"))
numeric_columns = set(schema["numeric_columns"])
date_columns = set(schema["date_columns"])
percent_columns = set(schema["percent_columns"])

workbook = xlsxwriter.Workbook(
    OUTPUT,
    {
        "constant_memory": True,
        "strings_to_numbers": False,
        "strings_to_formulas": False,
        "strings_to_urls": False,
    },
)
workbook.set_properties(
    {
        "title": "DANO — общая таблица по клиентам",
        "subject": "Агрегация демографии, штрафов и топливных транзакций по client_id",
        "company": "DATA SWAGA team",
    }
)
sheet = workbook.add_worksheet("Клиенты")
sheet.hide_gridlines(2)
sheet.freeze_panes(1, 1)
sheet.set_zoom(80)
sheet.set_tab_color("#FACC15")
sheet.set_default_row(15)

header_base = {"bold": True, "font_color": "#FFFFFF", "font_name": "Arial", "font_size": 9,
               "align": "center", "valign": "vcenter", "text_wrap": True, "border": 1,
               "border_color": "#FFFFFF"}
header_formats = {
    "demo": workbook.add_format({**header_base, "bg_color": "#374151"}),
    "fines_2025": workbook.add_format({**header_base, "bg_color": "#4B5563"}),
    "fines_2026": workbook.add_format({**header_base, "bg_color": "#B45309"}),
    "fuel": workbook.add_format({**header_base, "bg_color": "#0F766E"}),
}
text_format = workbook.add_format({"font_name": "Arial", "font_size": 9, "valign": "vcenter"})
integer_format = workbook.add_format({"font_name": "Arial", "font_size": 9, "num_format": "#,##0", "valign": "vcenter"})
decimal_format = workbook.add_format({"font_name": "Arial", "font_size": 9, "num_format": "#,##0.00", "valign": "vcenter"})
percent_format = workbook.add_format({"font_name": "Arial", "font_size": 9, "num_format": "0.0%", "valign": "vcenter"})
date_format = workbook.add_format({"font_name": "Arial", "font_size": 9, "num_format": "yyyy-mm-dd", "valign": "vcenter"})


def section_for(header: str) -> str:
    if header.startswith("fuel_"):
        return "fuel"
    if header.startswith("fines_2026") or header.startswith("fine_amount_2026") or header.startswith("severe_fines_2026") or header.startswith("offence_") or header.startswith("top_offence") or header.startswith("fines_change"):
        return "fines_2026"
    if header.startswith("fines_2025") or header.startswith("fines_last"):
        return "fines_2025"
    return "demo"


def column_width(header: str) -> int:
    if header == "client_id":
        return 29
    if header == "auto_document_ids":
        return 48
    if header == "top_offence_2026":
        return 48
    if header in {"registration_region", "auto_marks", "auto_colors", "top_offence_region_2026"}:
        return 24
    if header == "subscription_creation_date":
        return 16
    return 15


with INPUT.open("r", encoding="utf-8", newline="") as stream:
    reader = csv.reader(stream, delimiter="\t")
    headers = next(reader)
    for col, header in enumerate(headers):
        sheet.write(0, col, header, header_formats[section_for(header)])
        sheet.set_column(col, col, column_width(header))
    sheet.set_row(0, 72)

    last_row = 0
    for row_idx, row in enumerate(reader, start=1):
        last_row = row_idx
        for col_idx, header in enumerate(headers):
            value = row[col_idx] if col_idx < len(row) else ""
            if value == "":
                sheet.write_blank(row_idx, col_idx, None, text_format)
            elif header in date_columns:
                sheet.write_datetime(row_idx, col_idx, datetime.strptime(value, "%Y-%m-%d"), date_format)
            elif header in numeric_columns:
                number = float(value)
                if header in percent_columns:
                    sheet.write_number(row_idx, col_idx, number, percent_format)
                elif any(token in header for token in ("price", "amount", "volume", "spend", "engine_", "vehicle_year")):
                    sheet.write_number(row_idx, col_idx, number, decimal_format)
                else:
                    sheet.write_number(row_idx, col_idx, number, integer_format)
            else:
                sheet.write_string(row_idx, col_idx, value, text_format)

sheet.autofilter(0, 0, last_row, len(headers) - 1)
pct_col = headers.index("fines_change_pct_2026_vs_2025")
sheet.conditional_format(
    1,
    pct_col,
    last_row,
    pct_col,
    {
        "type": "3_color_scale",
        "min_color": "#FCA5A5",
        "mid_type": "num",
        "mid_value": 0,
        "mid_color": "#FEF3C7",
        "max_color": "#86EFAC",
    },
)
workbook.close()

# Reopen the saved file and validate its physical structure and key workbook features.
saved = load_workbook(OUTPUT, read_only=False, data_only=False)
assert saved.sheetnames == ["Клиенты"]
ws = saved["Клиенты"]
assert ws.max_row == schema["checks"]["rows"] + 1
assert ws.max_column == schema["checks"]["columns"]
assert ws.freeze_panes == "B2"
assert ws.auto_filter.ref == f"A1:{ws.cell(1, ws.max_column).column_letter}{ws.max_row}"
assert ws.cell(1, 1).value == "client_id"
assert ws.cell(2, 1).value is not None
assert ws.cell(ws.max_row, 1).value is not None
saved.close()

# Small visual-QA copy derived from the final workbook; not a deliverable.
preview_book = load_workbook(OUTPUT, read_only=False, data_only=False)
preview_sheet = preview_book["Клиенты"]
preview_sheet.delete_rows(23, preview_sheet.max_row - 22)
preview_sheet.delete_cols(15, preview_sheet.max_column - 14)
preview_sheet.sheet_properties.pageSetUpPr.fitToPage = True
preview_sheet.page_setup.orientation = "landscape"
preview_sheet.page_setup.fitToWidth = 1
preview_sheet.page_setup.fitToHeight = 1
preview_sheet.print_area = "A1:N22"
preview_book.save(ROOT / "tmp" / "client_master" / "final_preview.xlsx")
preview_book.close()

print(
    json.dumps(
        {
            "output": str(OUTPUT),
            "rows": schema["checks"]["rows"],
            "columns": schema["checks"]["columns"],
            "unique_client_ids": schema["checks"]["unique_client_ids"],
            "fines_2025_total": schema["checks"]["fines_2025_total"],
            "fines_2026_total": schema["checks"]["fines_2026_total"],
            "fuel_transactions_total": schema["checks"]["fuel_transactions_total"],
        },
        ensure_ascii=False,
        indent=2,
    )
)
